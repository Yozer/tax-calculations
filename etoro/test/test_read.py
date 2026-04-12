import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))

import pytest
from datetime import datetime
from decimal import Decimal
from unittest.mock import patch, MagicMock
from openpyxl import Workbook

import calculate_tax as ct


def create_test_workbook(account_activity_rows, closed_positions_rows,
                         dividends_rows=None, summary_rows=None):
    """Create an in-memory workbook with test data."""
    wb = Workbook()

    # Account Activity sheet
    ws_aa = wb.active
    ws_aa.title = 'Account Activity'
    aa_headers = ['Position ID', 'Type', 'Amount', 'Date', 'Details', 'Asset type', 'Realized Equity Change']
    ws_aa.append(aa_headers)
    for row in account_activity_rows:
        ws_aa.append(row)

    # Closed Positions sheet
    ws_cp = wb.create_sheet('Closed Positions')
    cp_headers = ['Position ID', 'Action', 'Amount', 'Open Date', 'Close Date', 'Type']
    ws_cp.append(cp_headers)
    for row in closed_positions_rows:
        ws_cp.append(row)

    # Dividends sheet
    ws_div = wb.create_sheet('Dividends')
    div_headers = ['Position ID', 'Instrument Name', 'Net Dividend Received (USD)',
                   'Withholding Tax Rate (%)', 'Withholding Tax Amount (USD)', 'Date of Payment']
    ws_div.append(div_headers)
    if dividends_rows:
        for row in dividends_rows:
            ws_div.append(row)

    # Financial Summary sheet
    ws_fs = wb.create_sheet('Financial Summary')
    fs_headers = ['Name', 'Amount\n in (USD)']
    ws_fs.append(fs_headers)
    if summary_rows:
        for row in summary_rows:
            ws_fs.append(row)

    return wb


# ============================================================
# read() - Transaction reading via mock workbook
# ============================================================

class TestRead:
    @patch('calculate_tax.load_workbook')
    def test_stock_position_closed(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Position closed', 150.0, '15/06/2025 10:00:00', 'AAPL/USD', 'Stocks', 50.0],
            ],
            closed_positions_rows=[
                ['100', 'Buy AAPL', 100.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'Real'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, grouped_txns, grouped_cp = ct.read('test.xlsx')

        assert len(entries) == 1
        e = entries[0]
        assert e['type'] == ct.StockType
        assert e['is_cfd'] is False
        assert e['open_amount'] == Decimal('100.0')
        assert e['close_amount'] == Decimal('150.0')
        assert e['equity_change'] == Decimal('50.0')

    @patch('calculate_tax.load_workbook')
    def test_crypto_open_and_close(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['200', 'Open Position', 500.0, '01/03/2025 10:00:00', 'BTC/USD', 'Crypto', 0],
                ['200', 'Position closed', 600.0, '15/06/2025 10:00:00', 'BTC/USD', 'Crypto', 100.0],
            ],
            closed_positions_rows=[
                ['200', 'Buy BTC', 500.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'Real'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 2
        # First: open position (crypto buy -> negative amount)
        assert entries[0]['type'] == ct.CryptoType
        assert entries[0]['amount'] == Decimal('-500.0')
        # Second: close position (crypto sell)
        assert entries[1]['type'] == ct.CryptoType
        assert entries[1]['amount'] == Decimal('600.0')

    @patch('calculate_tax.load_workbook')
    def test_cfd_position(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['300', 'Position closed', -50.0, '15/06/2025 10:00:00', 'OIL/USD', 'CFD', -50.0],
            ],
            closed_positions_rows=[
                ['300', 'Buy OIL', 200.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'CFD'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['is_cfd'] is True
        assert entries[0]['type'] == ct.StockType  # CFD asset type maps to StockType

    @patch('calculate_tax.load_workbook')
    def test_overnight_fee(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['400', 'Overnight fee', -2.5, '15/06/2025 10:00:00', 'daily', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['type'] == ct.FeeType
        assert entries[0]['amount'] == Decimal('-2.5')

    @patch('calculate_tax.load_workbook')
    def test_dividend_transaction(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['500', 'Dividend', 10.0, '15/06/2025 10:00:00', 'some dividend', 'Stocks', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['type'] == ct.DividendType

    @patch('calculate_tax.load_workbook')
    def test_interest_payment(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['600', 'Interest Payment', 3.0, '15/06/2025 10:00:00', '', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['type'] == ct.InterestType

    @patch('calculate_tax.load_workbook')
    def test_adjustment(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['700', 'Adjustment', 1.5, '15/06/2025 10:00:00', '', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['type'] == ct.AdjustmentType

    @patch('calculate_tax.load_workbook')
    def test_index_price_adjustment(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['750', 'Index price adjustment', 0.8, '15/06/2025 10:00:00', '', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['type'] == ct.IndexAdjustmentType

    @patch('calculate_tax.load_workbook')
    def test_withdrawal_conversion_fee_nonzero(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['800', 'Withdrawal Conversion Fee', -5.0, '15/06/2025 10:00:00', '', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['type'] == ct.FeeType
        assert entries[0]['amount'] == Decimal('-5.0')
        assert entries[0]['is_cfd'] is True

    @patch('calculate_tax.load_workbook')
    def test_withdrawal_conversion_fee_zero_skipped(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['801', 'Withdrawal Conversion Fee', 0, '15/06/2025 10:00:00', '', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 0

    @patch('calculate_tax.load_workbook')
    def test_withdraw_fee_zero_ok(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['900', 'Withdraw Fee', 0, '15/06/2025 10:00:00', '', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 0

    @patch('calculate_tax.load_workbook')
    def test_withdraw_fee_nonzero_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['901', 'Withdraw Fee', -5.0, '15/06/2025 10:00:00', '', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Unsupported withdraw fee'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_deposit_conversion_fee_nonzero_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['902', 'Deposit Conversion Fee', -2.0, '15/06/2025 10:00:00', '', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Unsupported withdraw fee'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_ignored_transactions_skipped(self, mock_lwb):
        rows = [[str(i), trans_type, 0, '15/06/2025 10:00:00', '', 'Stocks', 0]
                for i, trans_type in enumerate(ct.ignored_transactions)]
        wb = create_test_workbook(account_activity_rows=rows, closed_positions_rows=[])
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 0

    @patch('calculate_tax.load_workbook')
    def test_unknown_transaction_type_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['999', 'SomethingNew', 10.0, '15/06/2025 10:00:00', '', 'Stocks', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Unknown transaction type'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_wrong_year_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Deposit', 500.0, '15/06/2024 10:00:00', '', 'Stocks', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Invalid year'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_none_pos_id_skipped(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                [None, 'Deposit', 500.0, '15/06/2025 10:00:00', '', 'Stocks', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 0

    @patch('calculate_tax.load_workbook')
    def test_stock_open_position_non_crypto_skipped(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Open Position', 500.0, '01/03/2025 10:00:00', 'AAPL/USD', 'Stocks', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 0

    @patch('calculate_tax.load_workbook')
    def test_negative_crypto_buy_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Open Position', -100.0, '01/03/2025 10:00:00', 'BTC/USD', 'Crypto', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Negative crypto buy'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_sell_non_cfd_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Position closed', 150.0, '15/06/2025 10:00:00', 'AAPL/USD', 'Stocks', 50.0],
            ],
            closed_positions_rows=[
                ['100', 'Sell AAPL', 100.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'Real'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Not CFD that is sell'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_sell_cfd_ok(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Position closed', -50.0, '15/06/2025 10:00:00', 'OIL', 'CFD', -50.0],
            ],
            closed_positions_rows=[
                ['100', 'Sell OIL', 200.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'CFD'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['is_cfd'] is True

    @patch('calculate_tax.load_workbook')
    def test_multiple_closed_positions_same_id_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Position closed', 150.0, '15/06/2025 10:00:00', 'AAPL/USD', 'Stocks', 50.0],
            ],
            closed_positions_rows=[
                ['100', 'Buy AAPL', 100.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'Real'],
                ['100', 'Buy AAPL', 100.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'Real'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='More than one closed position'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_negative_open_amount_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Position closed', 150.0, '15/06/2025 10:00:00', 'AAPL/USD', 'Stocks', 50.0],
            ],
            closed_positions_rows=[
                ['100', 'Buy AAPL', -100.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'Real'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Negative open_amount'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_negative_amount_non_cfd_raises(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Position closed', -10.0, '15/06/2025 10:00:00', 'AAPL/USD', 'Stocks', -10.0],
            ],
            closed_positions_rows=[
                ['100', 'Buy AAPL', 100.0, '01/03/2025 10:00:00', '15/06/2025 10:00:00', 'Real'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Negative amount'):
            ct.read('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_t_plus_2_for_real_stocks(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Position closed', 150.0, '15/06/2025 10:00:00', 'AAPL/USD', 'Stocks', 50.0],
            ],
            closed_positions_rows=[
                # Thursday open, so T+2 = Monday
                ['100', 'Buy AAPL', 100.0, '12/06/2025 10:00:00', '15/06/2025 10:00:00', 'Real'],
            ],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), patch.object(ct, 'use_t_plus_2', True):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        # Open: Thursday Jun 12 + T2 = Monday Jun 16
        assert entries[0]['open_date'] == datetime(2025, 6, 16, 10, 0, 0)

    @patch('calculate_tax.load_workbook')
    def test_weekend_refund(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Weekend refund', 1.5, '15/06/2025 10:00:00', 'weekend fee', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['type'] == ct.RefundType

    @patch('calculate_tax.load_workbook')
    def test_overnight_refund(self, mock_lwb):
        wb = create_test_workbook(
            account_activity_rows=[
                ['100', 'Overnight refund', 0.5, '15/06/2025 10:00:00', 'daily', 'CFD', 0],
            ],
            closed_positions_rows=[],
        )
        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            entries, _, _ = ct.read('test.xlsx')

        assert len(entries) == 1
        assert entries[0]['type'] == ct.RefundType


# ============================================================
# read_summary
# ============================================================

class TestReadSummary:
    @patch('calculate_tax.load_workbook')
    def test_basic_summary(self, mock_lwb):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Financial Summary'
        ws.append(['Name', 'Amount\n in (USD)'])
        ws.append(['CFDs (Profit or Loss)', 100])
        ws.append(['Stocks (Profit or Loss)', 200])
        ws.append(['ETFs (Profit or Loss)', 50])
        ws.append(['Crypto (Profit or Loss)', 80])
        ws.append(['Stock and ETF Dividends (Profit)', 30])
        ws.append(['CFD Dividends (Profit or Loss)', 10])
        ws.append(['Fees (overnight, withdrawal, admin)', -15])
        ws.append(['SDRT Charge', -5])
        ws.append(['Income from Refunds', 3])
        ws.append(['Index adjustments', 2])
        ws.append(['Total Interest payments by eToro EU', 7])
        ws.append(['Total Return Swaps (Profit or Loss)', 0])
        ws.append(['Income from Airdrops', 0])
        ws.append(['Income from Staking', 0])
        ws.append(['Income from Corporate Actions', 0])
        ws.append(['Spread fee on CFDs', 25])
        ws.append(['Spread fee on crypto', 10])
        ws.append(['Spread fee on Total Return Swaps (TRS)', 0])
        ws.append(['Spread fee on stocks', 5])
        ws.append(['Spread fee on ETFs', 3])

        mock_lwb.return_value = wb
        stock, crypto, div, fees, interest, refunds, idx_adj = ct.read_summary('test.xlsx')

        assert stock == Decimal('350')   # 100 + 200 + 50
        assert crypto == Decimal('80')
        assert div == Decimal('40')      # 30 + 10
        assert fees == Decimal('-20')    # -15 + -5
        assert interest == Decimal('7')
        assert refunds == Decimal('3')
        assert idx_adj == Decimal('2')

    @patch('calculate_tax.load_workbook')
    def test_nonzero_unsupported_raises(self, mock_lwb):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Financial Summary'
        ws.append(['Name', 'Amount\n in (USD)'])
        ws.append(['Total Return Swaps (Profit or Loss)', 15])  # non-zero!

        mock_lwb.return_value = wb
        with pytest.raises(Exception, match='Unupported: non-zero'):
            ct.read_summary('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_unknown_row_raises(self, mock_lwb):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Financial Summary'
        ws.append(['Name', 'Amount\n in (USD)'])
        ws.append(['Some New Category', 100])

        mock_lwb.return_value = wb
        with pytest.raises(Exception, match='Unupported: unknown column'):
            ct.read_summary('test.xlsx')


# ============================================================
# read_dividend_taxes
# ============================================================

class TestReadDividendTaxes:
    @patch('calculate_tax.load_workbook')
    def test_basic_dividend_taxes(self, mock_lwb):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Dividends'
        ws.append(['Position ID', 'Instrument Name', 'Net Dividend Received (USD)',
                    'Withholding Tax Rate (%)', 'Withholding Tax Amount (USD)', 'Date of Payment'])
        ws.append(['100', 'Apple Inc', '8.50', '15%', '1.50', '15/06/2025'])

        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            dividend_taxes, raw = ct.read_dividend_taxes('test.xlsx')

        assert '100' in dividend_taxes
        assert len(dividend_taxes['100']) == 1
        dt = dividend_taxes['100'][0]
        assert dt['Net Dividend Received (USD)'] == Decimal('8.50')
        assert dt['Withholding Tax Rate (%)'] == Decimal('0.15')
        assert dt['Withholding Tax Amount (USD)'] == Decimal('1.50')
        assert dt['Date of Payment'] == datetime(2025, 6, 15)

    @patch('calculate_tax.load_workbook')
    def test_none_pos_id_skipped(self, mock_lwb):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Dividends'
        ws.append(['Position ID', 'Instrument Name', 'Net Dividend Received (USD)',
                    'Withholding Tax Rate (%)', 'Withholding Tax Amount (USD)', 'Date of Payment'])
        ws.append([None, 'Apple Inc', '8.50', '15%', '1.50', '15/06/2025'])

        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            dividend_taxes, raw = ct.read_dividend_taxes('test.xlsx')

        assert len(dividend_taxes) == 0

    @patch('calculate_tax.load_workbook')
    def test_wrong_year_raises(self, mock_lwb):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Dividends'
        ws.append(['Position ID', 'Instrument Name', 'Net Dividend Received (USD)',
                    'Withholding Tax Rate (%)', 'Withholding Tax Amount (USD)', 'Date of Payment'])
        ws.append(['100', 'Apple Inc', '8.50', '15%', '1.50', '15/06/2024'])

        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025), pytest.raises(Exception, match='Invalid year'):
            ct.read_dividend_taxes('test.xlsx')

    @patch('calculate_tax.load_workbook')
    def test_multiple_dividends_same_position(self, mock_lwb):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Dividends'
        ws.append(['Position ID', 'Instrument Name', 'Net Dividend Received (USD)',
                    'Withholding Tax Rate (%)', 'Withholding Tax Amount (USD)', 'Date of Payment'])
        ws.append(['100', 'Apple Inc', '8.50', '15%', '1.50', '15/06/2025'])
        ws.append(['100', 'Apple Inc', '5.00', '15%', '0.88', '20/09/2025'])

        mock_lwb.return_value = wb
        with patch.object(ct, 'year', 2025):
            dividend_taxes, raw = ct.read_dividend_taxes('test.xlsx')

        assert len(dividend_taxes['100']) == 2
